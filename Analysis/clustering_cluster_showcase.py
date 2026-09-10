"""
clustering_cluster_showcase.py — Detailed profile of the winning config's
clusters + a cross-config alignment showing how each cluster is recovered
(or not) by the other 5 (model, condition) configurations.

For the config chosen by clustering_model_selection.py:

  Table A — per-cluster detail:
      cluster, n, n_pct, year_mean, top_terms, stability, intra_cohesion,
      silhouette_cluster, domains, persistence (#configs with Jaccard ≥ 0.5)

  Table B — cross-config alignment matrix:
      rows   = winning config's clusters (with short TF-IDF label)
      cols   = the other 5 configs at the same resolution
      cells  = best-matching cluster ID + best-Jaccard

Both tables are saved as CSVs and as a publication DOCX.

Usage:
    python analysis/clustering_cluster_showcase.py
    python analysis/clustering_cluster_showcase.py --config pubmedbert__raw__r008
"""
from __future__ import annotations
import argparse, os, sys, io
os.environ["PYTHONIOENCODING"] = "utf-8"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pathlib import Path
from collections import Counter

import numpy as np
import pandas as pd

from _clustering_utils import (  # noqa: E402
    PROC, MODEL_TAGS, CONDITIONS, DOMAIN_RE,
    load_embeddings, load_papers, compute_domain_flags,
    top_tfidf_terms, best_jaccard, config_key, MIN_CLUSTER_SIZE,
)

STAB_DIR = PROC / "clustering_stability"
UMAP_DIR = STAB_DIR / "umap"
ROOT = PROC.parents[1]


def parse_config(cfg: str):
    parts = cfg.split("__")
    return parts[0], parts[1], int(parts[2].lstrip("r")) / 1000.0


def intra_cohesion(embs: np.ndarray, mask: np.ndarray,
                    sample_cap: int = 500, rng: np.random.RandomState = None) -> float:
    """Mean pairwise cosine within a cluster (sampled if > sample_cap)."""
    idx = np.where(mask)[0]
    if len(idx) < 2:
        return float("nan")
    if len(idx) > sample_cap:
        rng = rng or np.random.RandomState(42)
        idx = rng.choice(idx, size=sample_cap, replace=False)
    sub = embs[idx]
    sim = sub @ sub.T  # cosine (embs are L2-normalized)
    n = sim.shape[0]
    tri = sim[np.triu_indices(n, k=1)]
    return float(tri.mean())


def per_cluster_silhouette(coords: np.ndarray, memberships: np.ndarray,
                            big: list, sample_cap: int = 3000) -> dict:
    """Per-cluster mean silhouette on the UMAP projection."""
    from sklearn.metrics import silhouette_samples
    rng = np.random.RandomState(42)
    # Sample
    if len(coords) > sample_cap:
        idx = rng.choice(len(coords), size=sample_cap, replace=False)
        coords_s = coords[idx]
        mem_s = memberships[idx]
    else:
        coords_s = coords
        mem_s = memberships
    # Per-sample silhouette
    if len(np.unique(mem_s)) < 2:
        return {c: float("nan") for c in big}
    sil = silhouette_samples(coords_s, mem_s)
    out = {}
    for c in big:
        m = mem_s == c
        out[c] = float(sil[m].mean()) if m.any() else float("nan")
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=str, default=None,
                        help="Config to profile (default: winner of "
                             "model_selection_scores.csv)")
    parser.add_argument("--resolution", type=float, default=0.008)
    parser.add_argument("--min-size", type=int, default=MIN_CLUSTER_SIZE)
    args = parser.parse_args()

    # Determine winner if not supplied
    if args.config is None:
        scores = pd.read_csv(STAB_DIR / "model_selection_scores.csv")
        winner = scores.sort_values("composite", ascending=False).iloc[0]
        args.config = winner["config"]
        print(f"Winner from model_selection_scores.csv: {args.config} "
              f"(composite={winner['composite']:.3f})")
    else:
        print(f"Profiling requested config: {args.config}")

    model, cond, resolution = parse_config(args.config)
    rtag = f"r{int(resolution*1000):03d}"

    # ── Load prerequisites
    embs = load_embeddings(model, cond)
    papers = compute_domain_flags(load_papers())
    mem = np.load(STAB_DIR / f"memberships_{args.config}.npy")
    sizes = Counter(mem)
    big = sorted([c for c, n in sizes.items() if n >= args.min_size],
                 key=lambda c: -sizes[c])
    print(f"  config has {len(big)} clusters ≥ {args.min_size} papers "
          f"(of {len(sizes)} total)")

    # UMAP for silhouette
    coord_path = UMAP_DIR / f"umap_{model}__{cond}.npy"
    coords = np.load(coord_path) if coord_path.exists() else None
    if coords is None:
        print("  [warning] UMAP coords missing — silhouette will be NaN")

    # TF-IDF (only for big clusters)
    print("  computing TF-IDF terms...")
    tfidf = top_tfidf_terms(papers, mem, top_k=8, min_size=args.min_size)

    # Per-cluster stability (from stability_full.csv)
    stab_full = pd.read_csv(STAB_DIR / "stability_full.csv")
    stab_cfg = stab_full[stab_full["config"] == args.config].set_index("cluster")

    # Per-cluster silhouette
    sil = per_cluster_silhouette(coords, mem, big) if coords is not None else {}

    # Persistence: for each winning cluster, fraction of OTHER configs with a
    # match at Jaccard ≥ 0.5
    # Load all other memberships at same resolution
    print("  computing cross-config alignment...")
    other_configs = [f"{m}__{c}__{rtag}" for m in MODEL_TAGS for c in CONDITIONS
                     if f"{m}__{c}__{rtag}" != args.config]
    other_mem = {ok: np.load(STAB_DIR / f"memberships_{ok}.npy")
                 for ok in other_configs
                 if (STAB_DIR / f"memberships_{ok}.npy").exists()}

    # For each cluster: best Jaccard in each other config + best matching cluster ID
    alignment_rows = []
    # Precompute TF-IDF for other configs (so we can label matched clusters)
    print("  computing TF-IDF for comparator configs...")
    other_tfidf = {ok: top_tfidf_terms(papers, om, top_k=4, min_size=args.min_size)
                   for ok, om in other_mem.items()}

    detail_rows = []
    n_corpus = len(mem)
    for c in big:
        members_set = set(np.where(mem == c)[0].tolist())
        dom_pcts = {d: float(papers.loc[np.where(mem == c)[0],
                                         f"dom_{d}"].mean()) for d in DOMAIN_RE}
        primary = ", ".join(
            f"{d}:{int(v*100)}%" for d, v in sorted(dom_pcts.items(),
                                                     key=lambda kv: -kv[1])
            if v >= 0.2)[:80] or "—"

        # Intra-cohesion
        coh = intra_cohesion(embs, mem == c)

        # Persistence count
        persist = 0
        align = {}
        for ok, om in other_mem.items():
            # Find best-matching cluster in ok
            members_idx = list(members_set)
            cnt = Counter(om[members_idx])
            best_j, best_id = 0.0, None
            for cand, overlap in cnt.items():
                other_members = np.where(om == cand)[0]
                union = len(members_set | set(other_members.tolist()))
                j = overlap / union if union > 0 else 0
                if j > best_j:
                    best_j, best_id = j, int(cand)
            if best_j >= 0.5:
                persist += 1
            align[ok] = {"jaccard": float(best_j),
                         "cluster": best_id,
                         "other_terms": " ".join(
                             other_tfidf[ok].get(best_id, [])[:3])
                             if best_id is not None else ""}

        # Stability-summary values
        stab_mean = float(stab_cfg.loc[c, "stability_mean"]) \
            if c in stab_cfg.index else float("nan")

        detail_rows.append({
            "cluster": int(c),
            "n_papers": int(sizes[c]),
            "n_pct": round(100 * sizes[c] / n_corpus, 2),
            "year_mean": round(papers.loc[np.where(mem == c)[0], "year"]
                                   .mean(), 1),
            "top_terms": " | ".join(tfidf.get(int(c), [])),
            "stability_mean": round(stab_mean, 3),
            "intra_cohesion": round(coh, 3),
            "silhouette": round(sil.get(c, float("nan")), 3),
            "persistence_of_5": f"{persist}/5",
            "domains": primary,
        })

        row = {"cluster": int(c),
               "n_papers": int(sizes[c]),
               "top_terms_short": " ".join(tfidf.get(int(c), [])[:3])}
        for ok in other_configs:
            row[f"{ok}__jaccard"] = round(align.get(ok, {}).get("jaccard", 0.0), 3)
            row[f"{ok}__cluster"] = align.get(ok, {}).get("cluster")
            row[f"{ok}__terms"] = align.get(ok, {}).get("other_terms", "")
        alignment_rows.append(row)

    detail_df = pd.DataFrame(detail_rows).sort_values(
        "stability_mean", ascending=False)
    alignment_df = pd.DataFrame(alignment_rows)

    out_showcase = STAB_DIR / f"cluster_showcase_{args.config}.csv"
    out_align = STAB_DIR / f"cluster_alignment_{args.config}.csv"
    detail_df.to_csv(out_showcase, index=False)
    alignment_df.to_csv(out_align, index=False)
    print(f"  saved {out_showcase.name}")
    print(f"  saved {out_align.name}")

    # ─────── DOCX ─────────────────────────────────────────────────────────
    from docx import Document
    from docx.shared import Pt, Inches
    print("\n  building DOCX...")
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(14)
        section.page_height = Inches(11)

    def h(t, level=1):
        return doc.add_heading(t, level=level)

    def add_docx_table(df, cols, col_widths=None, font_size=7):
        tbl = doc.add_table(rows=1 + len(df), cols=len(cols))
        tbl.style = "Light Grid"
        for j, c in enumerate(cols):
            cell = tbl.rows[0].cells[j]
            cell.text = c
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(font_size)
        for i, (_, row) in enumerate(df.iterrows()):
            for j, c in enumerate(cols):
                val = row[c]
                if isinstance(val, float):
                    val = f"{val:.2f}" if c.endswith("jaccard") \
                                         or "stability" in c \
                                         or "cohesion" in c \
                                         or c == "silhouette" \
                                         else f"{val:g}"
                tbl.rows[i+1].cells[j].text = str(val)[:90]
                for run in tbl.rows[i+1].cells[j].paragraphs[0].runs:
                    run.font.size = Pt(font_size)

    h(f"Cluster showcase — {args.config}", level=1)
    doc.add_paragraph(
        f"Detailed profile of the {len(detail_df)} clusters "
        f"(≥{args.min_size} papers) produced by the winning configuration. "
        f"Metrics:  stability = mean best-Jaccard against the other 23 "
        f"configurations; intra_cohesion = mean pairwise cosine similarity "
        f"within the cluster (L2-normalized raw embeddings, sampled to 500 "
        f"members); silhouette = mean silhouette on the 2-D UMAP projection; "
        f"persistence_of_5 = number of the 5 same-resolution other configs "
        f"whose closest cluster overlaps this one at Jaccard ≥ 0.5 "
        f"(\"universal\" clusters hit 5/5)."
    )

    h("Table A — winner clusters, detailed profile", level=2)
    cols_a = ["cluster", "n_papers", "n_pct", "year_mean", "top_terms",
              "stability_mean", "intra_cohesion", "silhouette",
              "persistence_of_5", "domains"]
    add_docx_table(detail_df, cols_a)

    h("Table B — cross-config alignment (Jaccard vs best-matching cluster)", level=2)
    doc.add_paragraph(
        "For each cluster in the winning config, we report the best-matching "
        "cluster in each of the 5 other (model, condition) configurations at "
        "the same resolution. A large Jaccard (≥ 0.5) together with "
        "matching TF-IDF terms means the cluster is recovered; a small "
        "Jaccard means the cluster is specific to this configuration."
    )

    cols_b = ["cluster", "n_papers", "top_terms_short"]
    for ok in other_configs:
        short = ok.replace("__", "/")
        cols_b.append(f"{ok}__jaccard")
        cols_b.append(f"{ok}__terms")
    # Rename columns in-place for display
    rename_map = {c: c.replace("__", "/") for c in cols_b
                   if "jaccard" in c or c.endswith("terms")}
    disp_df = alignment_df[cols_b].rename(columns=rename_map)
    add_docx_table(disp_df, list(disp_df.columns), font_size=6)

    h("Table C — per-config summary at the working resolution", level=2)
    summary = []
    this_cfg = args.config
    stab_matrix = pd.read_csv(STAB_DIR / "stability_matrix.csv", index_col=0)
    for cfg in [this_cfg] + other_configs:
        if cfg not in stab_matrix.index:
            continue
        cfg_mem_path = STAB_DIR / f"memberships_{cfg}.npy"
        if not cfg_mem_path.exists():
            continue
        cm = np.load(cfg_mem_path)
        sz = Counter(cm)
        summary.append({
            "config": cfg.replace("__", "/"),
            "n_clusters_total": len(sz),
            "n_big": sum(1 for s in sz.values() if s >= args.min_size),
            "largest_cluster": max(sz.values()),
            "median_cluster": int(np.median(list(sz.values()))),
            "mean_stability": round(stab_matrix.loc[cfg].drop(cfg).mean(), 3),
        })
    add_docx_table(pd.DataFrame(summary),
                   ["config", "n_clusters_total", "n_big", "largest_cluster",
                    "median_cluster", "mean_stability"])

    out_docx = ROOT / "Paper" / f"clustering_cluster_showcase.docx"
    doc.save(str(out_docx))
    print(f"  saved {out_docx}")


if __name__ == "__main__":
    main()
