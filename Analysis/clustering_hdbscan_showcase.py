"""
clustering_hdbscan_showcase.py — per-cluster detailed profile + cross-
config alignment for the HDBSCAN winner. Mirror of
`clustering_cluster_showcase.py`.

Outputs:
    Data/processed/clustering_hdbscan/cluster_showcase_{config}.csv
    Data/processed/clustering_hdbscan/cluster_alignment_{config}.csv
    Paper/clustering_hdbscan_cluster_showcase.docx
"""
from __future__ import annotations
import argparse, os, sys, io
os.environ["PYTHONIOENCODING"] = "utf-8"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pathlib import Path
from collections import Counter
import numpy as np
import pandas as pd

from _clustering_utils import (  # noqa: E402
    PROC, DOMAIN_RE, load_embeddings, load_papers, compute_domain_flags,
    top_tfidf_terms, best_jaccard, MIN_CLUSTER_SIZE,
)
from _hdbscan_utils import (  # noqa: E402
    HDB_DIR, MODEL_TAGS, CONDITIONS, hdb_config_key,
)

ROOT = PROC.parents[1]
LEIDEN_UMAP_DIR = PROC / "clustering_stability" / "umap"


def parse_config(cfg: str):
    parts = cfg.split("__")
    return parts[0], parts[1], int(parts[2].lstrip("mcs"))


def intra_cohesion(embs, mask, sample_cap=500):
    idx = np.where(mask)[0]
    if len(idx) < 2:
        return float("nan")
    if len(idx) > sample_cap:
        idx = np.random.RandomState(42).choice(idx, size=sample_cap, replace=False)
    sub = embs[idx]
    sim = sub @ sub.T
    n = sim.shape[0]
    return float(sim[np.triu_indices(n, k=1)].mean())


def per_cluster_silhouette(coords, memberships, big, sample_cap=3000):
    from sklearn.metrics import silhouette_samples
    rng = np.random.RandomState(42)
    valid = memberships >= 0
    if valid.sum() < 50:
        return {c: float("nan") for c in big}
    coords = coords[valid]; memberships = memberships[valid]
    if len(coords) > sample_cap:
        idx = rng.choice(len(coords), size=sample_cap, replace=False)
        coords = coords[idx]; memberships = memberships[idx]
    if len(np.unique(memberships)) < 2:
        return {c: float("nan") for c in big}
    sil = silhouette_samples(coords, memberships)
    return {c: float(sil[memberships == c].mean())
            if (memberships == c).any() else float("nan") for c in big}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=str, default=None)
    parser.add_argument("--min-size", type=int, default=MIN_CLUSTER_SIZE)
    args = parser.parse_args()

    if args.config is None:
        scores = pd.read_csv(HDB_DIR / "model_selection_scores.csv")
        winner = scores.sort_values("composite", ascending=False).iloc[0]
        args.config = winner["config"]
        print(f"Winner: {args.config} (composite={winner['composite']:.3f})")

    model, cond, mcs = parse_config(args.config)
    embs = load_embeddings(model, cond)
    papers = compute_domain_flags(load_papers())
    mem = np.load(HDB_DIR / f"memberships_{args.config}.npy")
    sizes = Counter(mem)
    big = sorted([c for c, n in sizes.items() if c >= 0 and n >= args.min_size],
                 key=lambda c: -sizes[c])
    noise_frac = float((mem == -1).mean())
    print(f"  {len(big)} big clusters ≥ {args.min_size} papers "
          f"(of {len(sizes) - (1 if -1 in sizes else 0)} total, "
          f"noise = {noise_frac:.1%})")

    coords_path = LEIDEN_UMAP_DIR / f"umap_{model}__{cond}.npy"
    coords = np.load(coords_path) if coords_path.exists() else None

    tfidf = top_tfidf_terms(papers, mem, top_k=8, min_size=args.min_size)
    stab_full = pd.read_csv(HDB_DIR / "stability_full.csv")
    stab_cfg = stab_full[stab_full["config"] == args.config].set_index("cluster")
    sil = per_cluster_silhouette(coords, mem, big) if coords is not None else {}

    # Persistence vs other configs with same mcs
    other_configs = [hdb_config_key(m, c, mcs)
                     for m in MODEL_TAGS for c in CONDITIONS
                     if hdb_config_key(m, c, mcs) != args.config]
    other_mem = {ok: np.load(HDB_DIR / f"memberships_{ok}.npy")
                 for ok in other_configs
                 if (HDB_DIR / f"memberships_{ok}.npy").exists()}
    other_tfidf = {ok: top_tfidf_terms(papers, om, top_k=4, min_size=args.min_size)
                   for ok, om in other_mem.items()}

    alignment_rows = []
    detail_rows = []
    n_corpus = len(mem)
    for c in big:
        members_set = set(np.where(mem == c)[0].tolist())
        sub = papers.iloc[list(members_set)]
        dom_pcts = {d: float(sub[f"dom_{d}"].mean()) for d in DOMAIN_RE}
        primary = ", ".join(
            f"{d}:{int(v*100)}%" for d, v in sorted(dom_pcts.items(),
                                                      key=lambda kv: -kv[1])
            if v >= 0.2)[:80] or "—"
        coh = intra_cohesion(embs, mem == c)

        persist = 0; align = {}
        for ok, om in other_mem.items():
            members_idx = list(members_set)
            cnt = Counter(om[members_idx])
            # drop noise from consideration
            cnt.pop(-1, None)
            best_j, best_id = 0.0, None
            for cand, overlap in cnt.items():
                other_members = np.where(om == cand)[0]
                union = len(members_set | set(other_members.tolist()))
                j = overlap / union if union > 0 else 0
                if j > best_j:
                    best_j, best_id = j, int(cand)
            if best_j >= 0.5:
                persist += 1
            align[ok] = {
                "jaccard": float(best_j),
                "cluster": best_id,
                "other_terms": " ".join(other_tfidf[ok].get(best_id, [])[:3])
                    if best_id is not None else "",
            }

        stab_mean = float(stab_cfg.loc[c, "stability_mean"]) \
            if c in stab_cfg.index else float("nan")
        detail_rows.append({
            "cluster": int(c),
            "n_papers": int(sizes[c]),
            "n_pct": round(100 * sizes[c] / n_corpus, 2),
            "year_mean": round(sub["year"].mean(), 1),
            "top_terms": " | ".join(tfidf.get(int(c), [])),
            "stability_mean": round(stab_mean, 3),
            "intra_cohesion": round(coh, 3),
            "silhouette": round(sil.get(c, float("nan")), 3),
            "persistence_of_others": f"{persist}/{len(other_mem)}",
            "domains": primary,
        })
        row = {"cluster": int(c), "n_papers": int(sizes[c]),
               "top_terms_short": " ".join(tfidf.get(int(c), [])[:3])}
        for ok in other_configs:
            row[f"{ok}__jaccard"] = round(align.get(ok, {}).get("jaccard", 0.0), 3)
            row[f"{ok}__cluster"] = align.get(ok, {}).get("cluster")
            row[f"{ok}__terms"]   = align.get(ok, {}).get("other_terms", "")
        alignment_rows.append(row)

    detail_df = pd.DataFrame(detail_rows).sort_values(
        "stability_mean", ascending=False)
    align_df = pd.DataFrame(alignment_rows)
    detail_df.to_csv(HDB_DIR / f"cluster_showcase_{args.config}.csv", index=False)
    align_df.to_csv(HDB_DIR / f"cluster_alignment_{args.config}.csv", index=False)

    # ── DOCX
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(14)
        section.page_height = Inches(11)

    def add_table(df, cols, font_size=7):
        tbl = doc.add_table(rows=1 + len(df), cols=len(cols))
        tbl.style = "Light Grid"
        for j, c in enumerate(cols):
            tbl.rows[0].cells[j].text = c
            for run in tbl.rows[0].cells[j].paragraphs[0].runs:
                run.bold = True; run.font.size = Pt(font_size)
        for i, (_, row) in enumerate(df.iterrows()):
            for j, c in enumerate(cols):
                val = row[c]
                if isinstance(val, float):
                    val = f"{val:.2f}" if c.endswith("jaccard") \
                                         or "stability" in c \
                                         or "cohesion" in c \
                                         or c == "silhouette" \
                                         else f"{val:g}"
                tbl.rows[i+1].cells[j].text = str(val)[:90]
                for run in tbl.rows[i+1].cells[j].paragraphs[0].runs:
                    run.font.size = Pt(font_size)

    doc.add_heading(f"HDBSCAN cluster showcase — {args.config}", level=1)
    doc.add_paragraph(
        f"Per-cluster profile of the {len(detail_df)} clusters "
        f"(≥ {args.min_size} papers) produced by the HDBSCAN winner "
        f"(model = {model}, condition = {cond}, min_cluster_size = {mcs}; "
        f"noise = {noise_frac:.1%}). "
        f"Stability columns mirror those from the Leiden showcase so the "
        f"two methods are directly comparable."
    )
    doc.add_heading("Table A — winner clusters, detailed profile", level=2)
    cols_a = ["cluster", "n_papers", "n_pct", "year_mean", "top_terms",
              "stability_mean", "intra_cohesion", "silhouette",
              "persistence_of_others", "domains"]
    add_table(detail_df, cols_a)

    doc.add_heading("Table B — cross-config alignment (Jaccard vs best-match)",
                     level=2)
    doc.add_paragraph(
        f"For each winner cluster, the best-matching cluster in each of the "
        f"other (model, condition) HDBSCAN configs at mcs={mcs}. Large "
        f"Jaccard plus matching terms ⇒ the cluster is recovered. Small "
        f"Jaccard ⇒ it is specific to the winner."
    )
    cols_b = ["cluster", "n_papers", "top_terms_short"]
    for ok in other_configs:
        cols_b += [f"{ok}__jaccard", f"{ok}__terms"]
    rename = {c: c.replace("__", "/") for c in cols_b if "jaccard" in c or c.endswith("terms")}
    add_table(align_df[cols_b].rename(columns=rename),
              [rename.get(c, c) for c in cols_b], font_size=6)

    out = ROOT / "Paper" / "clustering_hdbscan_cluster_showcase.docx"
    doc.save(str(out))
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
