"""
summarize_influential_nonreviews.py — Readable summary of the
experimental/modelling (non-review) selection. Parallel to
`summarize_influential_reviews.py`.

Emits:
    Paper/influential_nonreviews_summary.docx
        methods note, coverage by decade × full-text source, per-decade
        top-3 picks, and two supplementary rankings (top-100 by global
        cites, top-100 by within-corpus cites).

    Data/processed/influential_nonreviews_by_global_citations.csv
    Data/processed/influential_nonreviews_by_internal_citations.csv

Usage:
    python analysis/summarize_influential_nonreviews.py
    python analysis/summarize_influential_nonreviews.py --top-n 100
"""
from __future__ import annotations
import argparse, os, sys, io
os.environ["PYTHONIOENCODING"] = "utf-8"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from pathlib import Path
import pandas as pd

from _clustering_utils import PROC

ROOT = PROC.parents[1]
OUT_DOCX = ROOT / "Paper" / "influential_nonreviews_summary.docx"


def truncate(s: str, n: int = 100) -> str:
    s = str(s or "")
    return s if len(s) <= n else s[: n - 1] + "…"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--top-n", type=int, default=100,
                        help="Top-N for citation-only tables in the DOCX")
    args = parser.parse_args()

    all_path = PROC / "influential_nonreviews.csv"
    sel_path = PROC / "influential_nonreviews_selected.csv"
    if not (all_path.exists() and sel_path.exists()):
        raise SystemExit(
            "Missing manifests — run select_influential_nonreviews.py first."
        )
    all_df = pd.read_csv(all_path, low_memory=False)
    sel_df = pd.read_csv(sel_path, low_memory=False)

    # Coverage
    sel_df["decade"] = (sel_df["year"] // 10) * 10
    coverage = (sel_df.groupby(["decade", "fulltext_source"]).size()
                      .unstack(fill_value=0))
    for c in ["pmc", "libkey", "abstract_only"]:
        if c not in coverage.columns:
            coverage[c] = 0
    coverage = coverage[["pmc", "libkey", "abstract_only"]]
    coverage["total"] = coverage.sum(axis=1)
    coverage["pct_fulltext"] = (
        (coverage["pmc"] + coverage["libkey"]) / coverage["total"] * 100
    ).round(0).astype(int)
    print(coverage)

    # Rankings by raw citations
    by_global = (all_df.sort_values("cited_by_count", ascending=False)
                        .reset_index(drop=True))
    by_global["rank_global"] = by_global.index + 1
    by_global.to_csv(
        PROC / "influential_nonreviews_by_global_citations.csv", index=False)

    by_internal = (all_df.sort_values("within_corpus_cite", ascending=False)
                          .reset_index(drop=True))
    by_internal["rank_internal"] = by_internal.index + 1
    by_internal.to_csv(
        PROC / "influential_nonreviews_by_internal_citations.csv", index=False)

    # DOCX
    from docx import Document
    from docx.shared import Pt, Inches
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)

    def h(txt, level=1):
        doc.add_heading(txt, level=level)

    def para(txt, bold=False, size=11):
        p = doc.add_paragraph()
        run = p.add_run(txt); run.font.size = Pt(size); run.bold = bold
        return p

    def bullet(items):
        for it in items:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(it).font.size = Pt(10)

    def add_table(df, cols, font_size=8, max_col_chars=100):
        tbl = doc.add_table(rows=1 + len(df), cols=len(cols))
        tbl.style = "Light Grid"
        for j, c in enumerate(cols):
            tbl.rows[0].cells[j].text = str(c)
            for run in tbl.rows[0].cells[j].paragraphs[0].runs:
                run.bold = True; run.font.size = Pt(font_size)
        for i, (_, row) in enumerate(df.iterrows()):
            for j, c in enumerate(cols):
                val = row[c] if c in row.index else ""
                if isinstance(val, float):
                    if pd.isna(val):
                        val = ""
                    elif c in {"composite", "z_global", "z_internal",
                                "per_year_accrual"}:
                        val = f"{val:.2f}"
                    else:
                        val = f"{val:g}"
                tbl.rows[i + 1].cells[j].text = truncate(str(val), max_col_chars)
                for run in tbl.rows[i + 1].cells[j].paragraphs[0].runs:
                    run.font.size = Pt(font_size)

    h("Influential thalamus research papers  (non-review)", level=1)
    para(
        "Companion to the reviews summary — same scoring pipeline, same "
        "citation metrics — but filtered to experimental, modelling and "
        "clinical research papers. This is the \"primary literature\" "
        "counterpart: every PubMed record in the corpus that is NOT "
        "flagged as Review, Systematic Review, Meta-Analysis, "
        "Editorial, Comment, Letter, News, Erratum, Retraction, "
        "Biography, or Historical Article, and whose title does not "
        "match /perspective|opinion|hypothesis|viewpoint|essay|"
        "commentary/i."
    )
    para(f"Total candidates: {len(all_df):,}.  "
         f"Selected top-3 per year (1950-2026): {len(sel_df):,}.", bold=True)

    h("Global vs internal (within-corpus) citations", level=2)
    para(
        "Same definitions as in the reviews summary: global citations "
        "= OpenAlex `cited_by_count` (everywhere in the literature); "
        "internal citations = in-degree in the thalamus citation graph "
        "(how many of the ~18,900 papers in our own corpus cite this "
        "paper). The composite score averages z(log1p(global)) and "
        "z(within-corpus)."
    )

    h("Coverage by decade × full-text source", level=2)
    para(
        "Same full-text provenance: PMC (OA), institutional LibKey, "
        "or abstract-only fallback."
    )
    cov_disp = coverage.reset_index().rename(columns={"decade": "Decade"})
    add_table(cov_disp, ["Decade", "pmc", "libkey", "abstract_only",
                           "total", "pct_fulltext"])

    # Per-decade top-3 (research papers)
    h("Top-3 per year, organised by decade", level=2)
    para(
        "Per-year top-3 by composite score. Same organisation as the "
        "reviews-summary — useful when you want balanced per-era "
        "coverage of the primary literature for LLM frame analysis."
    )
    disp_cols = ["year", "pmid", "journal", "title", "cited_by_count",
                 "within_corpus_cite", "composite", "fulltext_source"]
    for dec, group in sel_df.sort_values(["decade", "composite"],
                                           ascending=[True, False]).groupby("decade"):
        h(f"{int(dec)}s  ({len(group)} papers)", level=3)
        g = group.sort_values(["year", "composite"], ascending=[True, False])
        add_table(g, disp_cols, max_col_chars=90)

    # Top-N by global
    h(f"Top-{args.top_n} research papers by global citations (OpenAlex)",
       level=2)
    para(
        "Ranked purely by raw `cited_by_count`. This is the "
        "\"thalamus hall of fame\" by overall scientific impact — "
        "the individual experimental / modelling papers that the "
        "whole of science cites the most."
    )
    top_g_cols = ["rank_global", "year", "pmid", "journal", "title",
                  "cited_by_count", "within_corpus_cite",
                  "per_year_accrual", "fulltext_source"]
    add_table(by_global.head(args.top_n), top_g_cols, max_col_chars=90)

    h(f"Top-{args.top_n} research papers by within-corpus citations",
       level=2)
    para(
        "Ranked by how often the paper is cited by other thalamus "
        "research. These are the primary-literature \"landmarks\" "
        "inside our subfield — the experimental and modelling "
        "touchstones that shape subsequent thalamus work."
    )
    top_i_cols = ["rank_internal", "year", "pmid", "journal", "title",
                  "within_corpus_cite", "cited_by_count",
                  "fulltext_source"]
    add_table(by_internal.head(args.top_n), top_i_cols, max_col_chars=90)

    doc.save(str(OUT_DOCX))
    print(f"\nSaved: {OUT_DOCX}")

    print("\nTop 15 by GLOBAL citations:")
    for _, r in by_global.head(15).iterrows():
        print(f"  #{int(r['rank_global']):>3} [{int(r['year'])}, "
              f"cites={int(r['cited_by_count'])}, "
              f"internal={int(r['within_corpus_cite'])}]  "
              f"{truncate(r['title'], 85)}")
    print("\nTop 15 by INTERNAL citations:")
    for _, r in by_internal.head(15).iterrows():
        print(f"  #{int(r['rank_internal']):>3} [{int(r['year'])}, "
              f"internal={int(r['within_corpus_cite'])}, "
              f"cites={int(r['cited_by_count'])}]  "
              f"{truncate(r['title'], 85)}")


if __name__ == "__main__":
    main()
